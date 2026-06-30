# -*- coding: utf-8 -*-
"""Insert generated charts into existing Word report at correct positions."""
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os, re

REPORT = r'C:\Users\id_30\Desktop\风险哨兵智能体_比亚迪风险应对决策支持系统.docx'
IMGDIR = r'C:\Users\id_30\Desktop\报告截图'

doc = Document(REPORT)

# Chart placement rules: find text markers and insert image after
placements = {
    '图A_蒙特卡洛VaR损失分布.png': [
        'VaR损失分布', '蒙特卡洛VaR', '5.1', 'Sobol序列生成',
    ],
    '图B_NSGA-II帕累托前沿与三方案.png': [
        '帕累托前沿', '5.2', 'NSGA-II优化', '22个解',
    ],
    '图C_三套策略多维对比.png': [
        '三套策略', '4.1', '对比', '多维',
    ],
}

# Strategy: find paragraphs containing key text and insert image after them
inserted = set()

for para in doc.paragraphs:
    text = para.text.strip()
    if not text:
        continue
    
    for fname, keywords in placements.items():
        if fname in inserted:
            continue
        img_path = os.path.join(IMGDIR, fname)
        if not os.path.exists(img_path):
            continue
        
        match_count = sum(1 for kw in keywords if kw in text)
        if match_count >= 2:
            # Insert image in a new paragraph after this one
            parent = para._element.getparent()
            idx = list(parent).index(para._element)
            
            # Determine which chapter for proper figure numbering
            if '5.1' in text or '蒙特卡洛' in text or 'Sobol' in text:
                fig_num = '图A'
                caption = f'\\n图A  蒙特卡洛VaR损失分布（Sobol序列10,000次模拟，比亚迪USD敞口480亿，年化波动率6.5%）'
            elif '5.2' in text or '帕累托' in text or 'NSGA' in text:
                fig_num = '图B'
                caption = f'\\n图B  NSGA-II帕累托前沿与三方案提取（47代收敛，22个帕累托解，P5/P50/P95分位）'
            elif '三套策略' in text or '4.1' in text or '多维' in text:
                fig_num = '图C'
                caption = f'\\n图C  三套策略多维对比（保守/平衡/激进的成本、剩余风险与对冲工具组成）'
            else:
                continue
            
            # Create new paragraph for image
            from docx.oxml.ns import qn
            new_para = doc.add_paragraph()
            # Move it right after the matched paragraph
            parent.insert(idx + 1, new_para._element)
            
            run = new_para.add_run()
            run.add_picture(img_path, width=Inches(5.5))
            new_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Caption below
            cap_para = doc.add_paragraph()
            parent.insert(idx + 2, cap_para._element)
            cap_run = cap_para.add_run(caption)
            cap_run.font.size = Pt(9)
            cap_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            inserted.add(fname)
            print(f'Inserted {fname} -> {fig_num}')

doc.save(REPORT.replace('.docx', '_含图表.docx'))
print(f'\nDone! Saved to: {REPORT.replace(".docx", "_含图表.docx")}')
print(f'Inserted {len(inserted)} charts')
